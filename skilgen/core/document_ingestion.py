from __future__ import annotations

import csv
import html
import json
import re
import tomllib
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

try:  # pragma: no cover - optional dependency
    import yaml
except ImportError:  # pragma: no cover - optional dependency
    yaml = None
try:  # pragma: no cover - optional dependency
    from bs4 import BeautifulSoup
except ImportError:  # pragma: no cover - optional dependency
    BeautifulSoup = None
try:  # pragma: no cover - optional dependency
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - optional dependency
    load_workbook = None
try:  # pragma: no cover - optional dependency
    from pypdf import PdfReader
except ImportError:  # pragma: no cover - optional dependency
    PdfReader = None
try:  # pragma: no cover - optional dependency
    from pptx import Presentation
except ImportError:  # pragma: no cover - optional dependency
    Presentation = None


PLAIN_TEXT_EXTENSIONS = {".md", ".markdown", ".txt", ".rst", ".log"}
STRUCTURED_TEXT_EXTENSIONS = {".json", ".yaml", ".yml", ".xml", ".csv", ".tsv", ".toml", ".ini", ".cfg"}
RICH_DOCUMENT_EXTENSIONS = {".docx", ".pdf", ".pptx", ".html", ".htm", ".xlsx"}
SUPPORTED_DOCUMENT_EXTENSIONS = PLAIN_TEXT_EXTENSIONS | STRUCTURED_TEXT_EXTENSIONS | RICH_DOCUMENT_EXTENSIONS


def normalize_extracted_text(text: str) -> str:
    text = text.replace("\x00", " ")
    lines: list[str] = []
    previous = ""
    for raw in text.splitlines():
        line = " ".join(raw.split()).strip()
        if not line:
            if previous:
                lines.append("")
            previous = ""
            continue
        if line == previous:
            continue
        lines.append(line)
        previous = line
    cleaned = "\n".join(lines)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


def detect_document_type(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in PLAIN_TEXT_EXTENSIONS:
        return "text"
    if suffix in {".yaml", ".yml"}:
        return "yaml"
    if suffix == ".json":
        return "json"
    if suffix == ".xml":
        return "xml"
    if suffix in {".csv", ".tsv"}:
        return "csv"
    if suffix == ".toml":
        return "toml"
    if suffix in {".ini", ".cfg"}:
        return "config"
    if suffix == ".docx":
        return "docx"
    if suffix == ".pdf":
        return "pdf"
    if suffix == ".pptx":
        return "pptx"
    if suffix in {".html", ".htm"}:
        return "html"
    if suffix == ".xlsx":
        return "xlsx"
    return suffix.lstrip(".") or "text"


def _flatten(value: object, *, prefix: str = "") -> list[str]:
    if value is None:
        return []
    if isinstance(value, dict):
        lines: list[str] = []
        for key, nested in value.items():
            next_prefix = f"{prefix}.{key}" if prefix else str(key)
            nested_lines = _flatten(nested, prefix=next_prefix)
            if nested_lines:
                lines.extend(nested_lines)
            else:
                lines.append(next_prefix)
        return lines
    if isinstance(value, list):
        lines: list[str] = []
        for index, nested in enumerate(value):
            next_prefix = f"{prefix}[{index}]" if prefix else f"[{index}]"
            nested_lines = _flatten(nested, prefix=next_prefix)
            if nested_lines:
                lines.extend(nested_lines)
            else:
                lines.append(next_prefix)
        return lines
    if isinstance(value, (str, int, float, bool)):
        rendered = str(value).strip()
        if not rendered:
            return []
        return [f"{prefix}: {rendered}" if prefix else rendered]
    return [f"{prefix}: {value}" if prefix else str(value)]


def _extract_docx(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    xml = re.sub(r"</w:p>", "\n", xml)
    xml = re.sub(r"<[^>]+>", "", xml)
    return html.unescape(xml)


def _extract_pdf(path: Path) -> str:
    if PdfReader is None:
        return ""
    reader = PdfReader(str(path))
    pages: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            pages.append(text)
    return "\n".join(pages)


def _extract_html(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    if BeautifulSoup is None:
        return re.sub(r"<[^>]+>", " ", raw)
    soup = BeautifulSoup(raw, "html.parser")
    return soup.get_text("\n")


def _extract_json(path: Path) -> str:
    payload = json.loads(path.read_text(encoding="utf-8"))
    return "\n".join(_flatten(payload))


def _extract_yaml(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    if yaml is None:
        return raw
    payload = yaml.safe_load(raw)
    return "\n".join(_flatten(payload))


def _extract_xml(path: Path) -> str:
    root = ET.fromstring(path.read_text(encoding="utf-8"))
    return "\n".join(part.strip() for part in root.itertext() if part.strip())


def _extract_csv(path: Path) -> str:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    rows: list[str] = []
    with path.open(encoding="utf-8", newline="") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        for row in reader:
            rendered = " | ".join(cell.strip() for cell in row if cell.strip())
            if rendered:
                rows.append(rendered)
    return "\n".join(rows)


def _extract_toml(path: Path) -> str:
    payload = tomllib.loads(path.read_text(encoding="utf-8"))
    return "\n".join(_flatten(payload))


def _extract_config(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def _extract_xlsx(path: Path) -> str:
    if load_workbook is None:
        return ""
    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        lines.append(f"[Sheet] {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = [str(value).strip() for value in row if value not in {None, ""}]
            if values:
                lines.append(" | ".join(values))
    return "\n".join(lines)


def _extract_pptx(path: Path) -> str:
    if Presentation is None:
        return ""
    presentation = Presentation(path)
    lines: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        lines.append(f"[Slide {index}]")
        for shape in slide.shapes:
            text = getattr(shape, "text", "").strip()
            if text:
                lines.append(text)
    return "\n".join(lines)


def extract_document_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in PLAIN_TEXT_EXTENSIONS:
        return normalize_extracted_text(path.read_text(encoding="utf-8", errors="ignore"))
    if suffix == ".docx":
        return normalize_extracted_text(_extract_docx(path))
    if suffix == ".pdf":
        return normalize_extracted_text(_extract_pdf(path))
    if suffix in {".html", ".htm"}:
        return normalize_extracted_text(_extract_html(path))
    if suffix == ".json":
        return normalize_extracted_text(_extract_json(path))
    if suffix in {".yaml", ".yml"}:
        return normalize_extracted_text(_extract_yaml(path))
    if suffix == ".xml":
        return normalize_extracted_text(_extract_xml(path))
    if suffix in {".csv", ".tsv"}:
        return normalize_extracted_text(_extract_csv(path))
    if suffix == ".toml":
        return normalize_extracted_text(_extract_toml(path))
    if suffix in {".ini", ".cfg"}:
        return normalize_extracted_text(_extract_config(path))
    if suffix == ".xlsx":
        return normalize_extracted_text(_extract_xlsx(path))
    if suffix == ".pptx":
        return normalize_extracted_text(_extract_pptx(path))
    return normalize_extracted_text(path.read_text(encoding="utf-8", errors="ignore"))
