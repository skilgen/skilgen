from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
import zipfile

try:
    from openpyxl import Workbook
except ImportError:  # pragma: no cover - optional dependency
    Workbook = None

try:
    from pptx import Presentation
except ImportError:  # pragma: no cover - optional dependency
    Presentation = None

from skilgen.core.document_ingestion import detect_document_type, extract_document_text, normalize_extracted_text
from skilgen.core.requirements import load_requirements


def _write_docx(path: Path, text: str) -> None:
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:body>"
        + "".join(f"<w:p><w:r><w:t>{part}</w:t></w:r></w:p>" for part in text.splitlines())
        + "</w:body></w:document>"
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        "</Types>"
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="word/document.xml"/>'
        "</Relationships>"
    )
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("word/document.xml", document_xml)


def _write_pdf(path: Path, text: str) -> None:
    objects = [
        "1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n",
        "2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n",
        "3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 144] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>\nendobj\n",
    ]
    stream = f"BT\n/F1 18 Tf\n36 96 Td\n({text}) Tj\nET\n"
    objects.append(f"4 0 obj\n<< /Length {len(stream.encode('latin-1'))} >>\nstream\n{stream}endstream\nendobj\n")
    objects.append("5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n")

    header = "%PDF-1.4\n"
    offsets: list[int] = []
    body = ""
    current = len(header.encode("latin-1"))
    for obj in objects:
        offsets.append(current)
        encoded = obj.encode("latin-1")
        body += obj
        current += len(encoded)
    xref_offset = current
    xref = ["xref\n", f"0 {len(objects) + 1}\n", "0000000000 65535 f \n"]
    xref.extend(f"{offset:010d} 00000 n \n" for offset in offsets)
    trailer = (
        "trailer\n"
        f"<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        "startxref\n"
        f"{xref_offset}\n"
        "%%EOF\n"
    )
    path.write_bytes((header + body + "".join(xref) + trailer).encode("latin-1"))


@unittest.skipUnless(Workbook is not None and Presentation is not None, "document ingestion extras are not installed")
class DocumentIngestionTests(unittest.TestCase):
    def test_extract_document_text_supports_multiple_formats(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)

            markdown = root / "requirements.md"
            markdown.write_text("Backend API endpoints\nFrontend components\n", encoding="utf-8")

            html_file = root / "spec.html"
            html_file.write_text("<html><body><h1>Architecture</h1><p>Use service boundaries.</p></body></html>", encoding="utf-8")

            json_file = root / "spec.json"
            json_file.write_text('{"service":"billing","owners":["platform","api"]}', encoding="utf-8")

            yaml_file = root / "spec.yaml"
            yaml_file.write_text("domains:\n  - backend\n  - frontend\n", encoding="utf-8")

            csv_file = root / "matrix.csv"
            csv_file.write_text("name,owner\nbilling,platform\n", encoding="utf-8")

            xml_file = root / "architecture.xml"
            xml_file.write_text("<root><domain>backend</domain><owner>platform</owner></root>", encoding="utf-8")

            docx_file = root / "requirements.docx"
            _write_docx(docx_file, "Backend services\nFrontend flows")

            pdf_file = root / "requirements.pdf"
            _write_pdf(pdf_file, "Backend PDF Requirements")

            xlsx_file = root / "skills.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Skills"
            sheet.append(["domain", "owner"])
            sheet.append(["backend", "platform"])
            workbook.save(xlsx_file)

            pptx_file = root / "architecture.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "Architecture"
            slide.placeholders[1].text = "Backend domain\nPlatform owner"
            presentation.save(pptx_file)

            self.assertIn("Backend API endpoints", extract_document_text(markdown))
            self.assertIn("Architecture", extract_document_text(html_file))
            self.assertIn("service: billing", extract_document_text(json_file))
            self.assertIn("domains[0]: backend", extract_document_text(yaml_file))
            self.assertIn("billing | platform", extract_document_text(csv_file))
            self.assertIn("backend", extract_document_text(xml_file))
            self.assertIn("Backend services", extract_document_text(docx_file))
            self.assertIn("Backend PDF Requirements", extract_document_text(pdf_file))
            self.assertIn("[Sheet] Skills", extract_document_text(xlsx_file))
            self.assertIn("Platform owner", extract_document_text(pptx_file))

            self.assertEqual(detect_document_type(pdf_file), "pdf")
            self.assertEqual(detect_document_type(docx_file), "docx")
            self.assertEqual(detect_document_type(yaml_file), "yaml")

    def test_load_requirements_supports_all_documented_formats(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)

            def make_source(path: Path) -> None:
                suffix = path.suffix.lower()
                if suffix in {".md", ".txt"}:
                    path.write_text("Backend API endpoints\nFrontend routes\n", encoding="utf-8")
                    return
                if suffix == ".docx":
                    _write_docx(path, "Backend API endpoints\nFrontend routes")
                    return
                if suffix == ".pdf":
                    _write_pdf(path, "Backend API endpoints and frontend routes")
                    return
                if suffix in {".html", ".htm"}:
                    path.write_text("<html><body><p>Backend API endpoints</p><p>Frontend routes</p></body></html>", encoding="utf-8")
                    return
                if suffix == ".json":
                    path.write_text('{"backend":"API endpoints","frontend":"routes"}', encoding="utf-8")
                    return
                if suffix in {".yaml", ".yml"}:
                    path.write_text("backend: API endpoints\nfrontend: routes\n", encoding="utf-8")
                    return
                if suffix == ".csv":
                    path.write_text("domain,detail\nbackend,API endpoints\nfrontend,routes\n", encoding="utf-8")
                    return
                if suffix == ".tsv":
                    path.write_text("domain\tdetail\nbackend\tAPI endpoints\nfrontend\troutes\n", encoding="utf-8")
                    return
                if suffix == ".xml":
                    path.write_text("<root><backend>API endpoints</backend><frontend>routes</frontend></root>", encoding="utf-8")
                    return
                if suffix == ".xlsx":
                    workbook = Workbook()
                    sheet = workbook.active
                    sheet.title = "Requirements"
                    sheet.append(["domain", "detail"])
                    sheet.append(["backend", "API endpoints"])
                    sheet.append(["frontend", "routes"])
                    workbook.save(path)
                    return
                if suffix == ".pptx":
                    presentation = Presentation()
                    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
                    slide.shapes.title.text = "Requirements"
                    slide.placeholders[1].text = "Backend API endpoints\nFrontend routes"
                    presentation.save(path)
                    return
                if suffix == ".toml":
                    path.write_text('backend = "API endpoints"\nfrontend = "routes"\n', encoding="utf-8")
                    return
                if suffix in {".ini", ".cfg"}:
                    path.write_text("[requirements]\nbackend = API endpoints\nfrontend = routes\n", encoding="utf-8")
                    return
                raise AssertionError(f"Unhandled suffix: {suffix}")

            paths = [
                root / "requirements.md",
                root / "requirements.txt",
                root / "requirements.docx",
                root / "requirements.pdf",
                root / "requirements.html",
                root / "requirements.htm",
                root / "requirements.json",
                root / "requirements.yaml",
                root / "requirements.yml",
                root / "requirements.csv",
                root / "requirements.tsv",
                root / "requirements.xml",
                root / "requirements.xlsx",
                root / "requirements.pptx",
                root / "requirements.toml",
                root / "requirements.ini",
                root / "requirements.cfg",
            ]
            for path in paths:
                make_source(path)

            for path in paths:
                with self.subTest(path=path.suffix.lower()):
                    context = load_requirements(path)
                    self.assertTrue(context.raw_text.strip())
                    self.assertTrue(context.summary)
                    self.assertTrue(context.domains["backend"])
                    self.assertTrue(context.domains["frontend"])

    def test_normalize_extracted_text_collapses_duplicate_noise(self) -> None:
        text = "Header\nHeader\n\n\nBackend API endpoints\n\x00Frontend routes\n"
        normalized = normalize_extracted_text(text)
        self.assertNotIn("\x00", normalized)
        self.assertIn("Backend API endpoints", normalized)
        self.assertIn("Frontend routes", normalized)


if __name__ == "__main__":
    unittest.main()
